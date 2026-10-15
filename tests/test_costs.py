"""Fees and costs (synthetic data, no network)."""
import io
from datetime import date, datetime

import pytest

from meerwaarde.engine.costs import costs
from meerwaarde.engine.portfolio import Portfolio
from meerwaarde.engine.tax import TaxCalculator, TaxRules
from meerwaarde.models import BUY, FEE, SELL, TOB, WITHDRAWAL, Transaction
from meerwaarde.prices import PriceBook
from meerwaarde.report import Report

TODAY = date(2026, 6, 30)
WORLD = "IE00B4L5Y983"          # in the fund catalog (TER known)


def tx(when, kind, asset, broker="A", **kw):
    return Transaction(broker, datetime.combine(when, datetime.min.time()), kind, asset, name=asset, **kw)


def make(txns, now):
    pf = Portfolio(txns)
    pb = PriceBook(fetch=False, today=TODAY)
    for a, v in now.items():
        pb.set_override(a, now=v)
    return Report(pf, pb, TaxCalculator(pf, pb, TaxRules(), today=TODAY), pf.warnings)


def test_costs_by_kind_and_year():
    from meerwaarde.funds import lookup
    txns = [tx(date(2025, 12, 31), BUY, WORLD, qty=100, gross_eur=10000, fees_eur=18.0, total_eur=-10018, asset_class="ETF"),
            tx(date(2025, 12, 31), TOB, WORLD, amount_eur=-12.0, asset_class="ETF"),     # info row, inside the 18
            tx(date(2026, 2, 1), FEE, "EUR", amount_eur=-2.5, asset_class="Cash/FX"),
            tx(date(2026, 1, 3), BUY, "BTC", broker="C", qty=1, gross_eur=100, fees_eur=0.25, total_eur=-100.25, asset_class="Crypto"),
            tx(date(2026, 3, 1), WITHDRAWAL, "BTC", broker="C", qty=-0.5, fee_qty=0.01, asset_class="Crypto"),
            tx(date(2026, 4, 1), SELL, "X", qty=-1, gross_eur=50, fees_eur=1.0, total_eur=49, asset_class="ETF"),
            tx(date(2026, 1, 1), BUY, "X", qty=1, gross_eur=40, fees_eur=1.0, total_eur=-41, asset_class="ETF")]
    c = costs(make(txns, {WORLD: 100.0, "BTC": 100.0}))
    got = {(r["year"], r["broker"], r["kind"]): r["amount"] for r in c.rows}
    assert got[(2025, "A", "Stock exchange tax")] == 12.0 and got[(2025, "A", "Commissions and conversion")] == 6.0
    assert got[(2026, "A", "Other fees")] == 2.5 and got[(2026, "C", "Commissions and conversion")] == 0.25
    assert got[(2026, "C", "Network fees")] == pytest.approx(0.01 * 100.25)
    ter = lookup(WORLD).ter
    assert got[(2026, "A", "Running costs (TER)")] == pytest.approx(6 * 100 * 100 * ter / 12, rel=1e-6)   # Jan–Jun at 100
    assert c.yearly_now["Running costs (TER)"] == pytest.approx(10000 * ter)
    assert c.yearly_now["Commissions and conversion"] == pytest.approx(6.0 + 0.25 + 2.0)
    assert "X" in c.unknown_ter[0] or c.unknown_ter == ["X"]


def test_costs_table_and_excel():
    from openpyxl import load_workbook
    from meerwaarde.export import to_xlsx
    txns = [tx(date(2026, 1, 2), BUY, WORLD, qty=10, gross_eur=1000, fees_eur=3.0, total_eur=-1003, asset_class="ETF")]
    rep = make(txns, {WORLD: 100.0})
    t = rep.costs_table(rep.costs())
    assert list(t.columns)[:1] == ["year"] and t.total.iloc[0] > 3.0
    ws = load_workbook(io.BytesIO(to_xlsx(rep)))["Costs"]
    assert ws["A1"].value == "Year"
    assert any(str(c.value).startswith("=SUM(") for row in ws.iter_rows() for c in row if c.value)
