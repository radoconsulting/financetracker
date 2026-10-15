"""Multi-year tax plan (synthetic data)."""
import io
from datetime import date, datetime

import pytest

from meerwaarde.engine.harvest import CostModel
from meerwaarde.engine.portfolio import Portfolio
from meerwaarde.engine.tax import TaxCalculator, TaxRules
from meerwaarde.engine.taxplan import plan
from meerwaarde.models import BUY, Transaction
from meerwaarde.prices import PriceBook
from meerwaarde.report import Report

TODAY = date(2026, 10, 14)
FREE = CostModel(spread=0.0, commission={"A": (0.0, 0.0)})


def make(txns, prices):
    pf = Portfolio(txns)
    pb = PriceBook(fetch=False, today=TODAY)
    for a, (now, ref) in prices.items():
        pb.set_override(a, now=now, reference=ref)
    return Report(pf, pb, TaxCalculator(pf, pb, TaxRules(), today=TODAY), pf.warnings)


def buy(asset, qty, price, when=datetime(2024, 1, 2), cls="ETF"):
    return Transaction("A", when, BUY, asset, name=asset, asset_class=cls, qty=qty, gross_eur=qty * price,
                       total_eur=-qty * price)


def test_uses_each_years_exemption_and_lowers_the_final_tax():
    rep = make([buy("X", 1000, 50.0)], {"X": (100.0, 60.0)})       # 40,000 of taxable gain today
    tp = plan(rep, growth=0.0, costs=FREE)
    assert [y.year for y in tp.years] == [2026, 2027, 2028, 2029, 2030]
    assert tp.years[0].room["you"] == pytest.approx(9000.0)          # 10,000 − 1,000 kept for the carry-forward
    assert tp.years[1].room["you"] == pytest.approx(10000.0)         # 10,000 + 1,000 carried − 1,000 kept
    assert sum(y.gain for y in tp.years) == pytest.approx(40000.0, abs=50)
    assert tp.tax_without == pytest.approx(0.10 * (40000 - 11000))  # 2031: own 10,000 + 1,000 carried
    assert tp.tax_with < 100 and tp.saving > 2800


def test_original_price_above_the_31dec_value_counts_before_2031():
    rep = make([buy("Y", 100, 120.0)], {"Y": (130.0, 100.0)})       # bought at 120, 31-12-2025 at 100, now 130
    tp = plan(rep, growth=0.0, costs=FREE)
    line = tp.years[0].lines[0]
    assert line.gain == pytest.approx(100 * 10.0)                   # gain from the original price (120), not 100
    assert line.window == pytest.approx(100 * 20.0)                 # 120 − 100 kept as purchase value
    assert tp.tax_without == 0.0                                    # 3,000 gain in 2031 fits in the exemption anyway


def test_costs_can_make_it_not_worth_it():
    rep = make([buy("Z", 10, 50.0)], {"Z": (60.0, 55.0)})          # 50 of gain
    tp = plan(rep, growth=0.0, costs=CostModel(commission={"A": (20.0, 0.0)}))
    assert all(not y.lines for y in tp.years) and tp.costs == 0.0


def test_partner_shares_split_the_room():
    from meerwaarde.profile import TaxProfile
    prof = TaxProfile()
    prof.set_partner("Partner")
    prof.household.brokers = {"A": {"you": 0.5, "partner": 0.5}}
    pf = Portfolio([buy("X", 1000, 50.0)])
    pb = PriceBook(fetch=False, today=TODAY)
    pb.set_override("X", now=100.0, reference=60.0)
    rep = Report(pf, pb, TaxCalculator(pf, pb, TaxRules(), today=TODAY), [], prof)
    tp = plan(rep, growth=0.0, costs=FREE)
    assert tp.years[0].gain == pytest.approx(18000.0, abs=50)       # 9,000 per person, half of each line each


def test_plan_sheet_and_cli_table():
    from openpyxl import load_workbook
    from meerwaarde.export import to_xlsx
    rep = make([buy("X", 1000, 50.0)], {"X": (100.0, 60.0)})
    assert not rep.tax_plan_table(rep.tax_plan(growth=0.0)).empty
    ws = load_workbook(io.BytesIO(to_xlsx(rep)))["Tax plan"]
    assert ws["A1"].value == "Year" and any(str(c.value).startswith("=D") for row in ws.iter_rows() for c in row if c.value)
