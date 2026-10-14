"""Dividend and interest tax, and the yearly tax report (synthetic data)."""
from datetime import date, datetime

import pytest

from meerwaarde.engine.household import Household, Person
from meerwaarde.engine.income_tax import IncomeTaxRules, income_tax
from meerwaarde.engine.portfolio import Portfolio
from meerwaarde.engine.tax import TaxCalculator, TaxRules
from meerwaarde.export import tax_report_xlsx
from meerwaarde.models import BUY, DIVIDEND, INTEREST, SELL, TOB, WITHHOLDING, IncomeRow, Transaction
from meerwaarde.pdf_report import tax_report_pdf
from meerwaarde.prices import PriceBook
from meerwaarde.profile import TaxProfile
from meerwaarde.report import Report
from meerwaarde.taxreport import tax_year_report

TODAY = date(2027, 2, 1)


def inc(kind, amount, cls="Stock", broker="B", asset="XS", name="Some Corp", year=2026):
    return IncomeRow(broker, datetime(year, 5, 1), asset, name, kind, amount, None, "EUR", "", cls)


def test_share_dividends_get_the_exemption_funds_and_interest_dont():
    rows = [inc(DIVIDEND, 1000.0), inc(WITHHOLDING, -150.0), inc(DIVIDEND, 400.0, cls="ETF", asset="FUND"),
            inc(INTEREST, 10.0, cls="Cash/FX", asset="EUR", name="Credit interest"),
            inc(WITHHOLDING, -2.0, cls="Cash/FX", asset="EUR", name="Withholding on credit interest"),
            inc(DIVIDEND, 999.0, year=2025)]
    (it,) = income_tax(rows, Household(), 2026).values()
    assert it.total("shares") == pytest.approx(850.0)                    # after 15% foreign tax
    assert it.exemption_used == pytest.approx(833.0)
    assert it.total("interest") == pytest.approx(8.0)
    assert it.taxable == pytest.approx(850 - 833 + 400 + 8)
    assert it.tax == pytest.approx(0.30 * (17 + 400 + 8))
    assert it.foreign_tax == pytest.approx(-152.0)


def test_partner_split_and_other_income():
    hh = Household([Person("you", "You"), Person("partner", "P")], brokers={"B": {"you": 0.5, "partner": 0.5}})
    rows = [inc(DIVIDEND, 2000.0)]
    res = income_tax(rows, hh, 2026, IncomeTaxRules(), other={"partner": {"interest": 100.0}})
    assert res["you"].exemption_used == pytest.approx(833.0) and res["partner"].exemption_used == pytest.approx(833.0)
    assert res["partner"].total("interest") == pytest.approx(100.0)
    assert res["you"].tax == pytest.approx(0.3 * (1000 - 833))


def tx(when, kind, asset, qty=0.0, total=None, gross=None, cls="ETF", broker="Alpha", **kw):
    return Transaction(broker=broker, ts=when, kind=kind, asset=asset, name=f"Fund {asset}", asset_class=cls,
                       qty=qty, total_eur=total, gross_eur=gross, fees_eur=0.0, source="x.pdf", **kw)


@pytest.fixture
def rep():
    txns = [
        tx(datetime(2024, 3, 1), BUY, "XX1", 100, total=-5000, gross=5000),
        tx(datetime(2026, 5, 4), SELL, "XX1", -20, total=1400, gross=1400),
        tx(datetime(2026, 5, 4), TOB, "XX1", 20, amount_eur=-1.68),
        tx(datetime(2026, 6, 1), DIVIDEND, "XX1", amount_eur=120.0),
        tx(datetime(2026, 7, 1), BUY, "SH1", 10, total=-1000, gross=1000, cls="Stock", broker="Beta"),
    ]
    pf = Portfolio(txns)
    pb = PriceBook(fetch=False, today=TODAY)
    pb.set_override("XX1", now=80.0, reference=60.0)
    pb.set_override("SH1", now=100.0)
    profile = TaxProfile()
    profile.year(2026).other_gains["you"] = {"crypto": -50.0}
    return Report(pf, pb, TaxCalculator(pf, pb, TaxRules(), today=TODAY), pf.warnings, profile)


def test_stock_exchange_tax_is_information_only(rep):
    assert rep.realized().pnl_eur.sum() == pytest.approx(1400 - 1000)      # the TOB row doesn't change profit
    t = rep.tob_summary(2026).set_index("broker")
    assert t.loc["Alpha", "amount_eur"] == pytest.approx(1.68) and t.loc["Alpha", "basis"] == "listed by the broker"
    assert t.loc["Beta", "amount_eur"] == pytest.approx(3.5)                # estimate: 0.35% of a share purchase


def test_yearly_report_and_carry_forward(rep):
    ytr = tax_year_report(rep, 2026)
    you = ytr.cgt["people"]["you"]
    assert you["net"] == pytest.approx(20 * (70 - 60) - 50)                 # sale from the 31-12-2025 price, plus the other loss
    assert ytr.exemption["you"].carry_out == 1000
    assert ytr.income["you"].tax == pytest.approx(36.0)
    assert ytr.complete
    rep.tax.rules.year = 2027
    rep.tax.rules.exemption = 10000
    assert rep.exemptions(2027) == {"you": 11000.0}
    pdf = tax_report_pdf(ytr)
    assert pdf[:4] == b"%PDF" and len(pdf) > 2000
    xlsx = tax_report_xlsx(ytr)
    assert xlsx[:2] == b"PK"


def test_report_workbook_formulas(rep):
    import io
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(tax_report_xlsx(tax_year_report(rep, 2026))))
    assert wb.sheetnames == ["Report", "Sales 2026", "Income 2026", "Stock exchange tax", "Exemption", "Checklist"]
    ws = wb["Report"]
    formulas = [c.value for row in ws.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("=")]
    assert any("$B$4" in f for f in formulas) and any("$B$5" in f for f in formulas)    # rates are input cells
