"""Reynders tax on bond funds (synthetic data)."""
from datetime import date, datetime

import pytest

from meerwaarde.engine.harvest import plan_harvest
from meerwaarde.engine.portfolio import Portfolio
from meerwaarde.engine.reynders import (ASSET_TEST, TIS, WORST_CASE, Reynders, ReyndersBook, ReyndersFund,
                                        ReyndersRules)
from meerwaarde.engine.tax import TaxCalculator, TaxRules
from meerwaarde.models import BUY, SELL, Lot, Transaction
from meerwaarde.prices import PriceBook
from meerwaarde.profile import TaxProfile
from meerwaarde.report import Report

TODAY = date(2026, 10, 14)
BOND = "XX0000000077"


def tx(when, kind, asset, qty, gross, cls="Bond ETF"):
    total = -gross if kind == BUY else gross
    return Transaction(broker="B", ts=when, kind=kind, asset=asset, name=f"Fund {asset}", asset_class=cls, qty=qty,
                       gross_eur=gross, fees_eur=0.0, total_eur=total, source="x.pdf")


def make(fund: ReyndersFund | None, txns=None, rules=None, profile=None, now=110.0, ref=100.0):
    txns = txns or [tx(datetime(2026, 2, 1), BUY, BOND, 100, 9000)]          # bought at 90 in 2026
    pf = Portfolio(txns)
    pb = PriceBook(fetch=False, today=TODAY)
    pb.set_override(BOND, now=now, reference=ref)
    book = ReyndersBook({BOND: fund} if fund else {})
    calc = TaxCalculator(pf, pb, TaxRules(), today=TODAY, reynders=Reynders(book, rules, pb.eur_per))
    return Report(pf, pb, calc, pf.warnings, profile or TaxProfile())


def lot(day, gross, qty=1.0):
    return Lot(datetime.combine(day, datetime.min.time()), qty, gross, gross)


def test_tis_method_uses_the_difference_between_sale_and_purchase():
    f = ReyndersFund(BOND, tis=[(date(2026, 1, 15), 2.0), (date(2026, 10, 1), 5.0)])
    r = Reynders(ReyndersBook({BOND: f}))
    assert r.unit_base(BOND, lot(date(2026, 2, 1), 90.0), 110.0, TODAY) == (3.0, TIS)


def test_tis_in_another_currency_is_converted_at_each_date():
    f = ReyndersFund(BOND, currency="USD", tis=[(date(2026, 1, 15), 2.0), (date(2026, 10, 1), 5.0)])
    rates = {date(2026, 1, 15): 0.9, date(2026, 10, 1): 0.8}
    r = Reynders(ReyndersBook({BOND: f}), eur_per=lambda cur, on: rates[on])
    base, method = r.unit_base(BOND, lot(date(2026, 2, 1), 90.0), 110.0, TODAY)
    assert base == pytest.approx(5.0 * 0.8 - 2.0 * 0.9) and method == TIS


def test_asset_test_when_there_is_no_tis_and_worst_case_without_anything():
    r = Reynders(ReyndersBook({BOND: ReyndersFund(BOND, debt_pct=0.6)}))
    assert r.unit_base(BOND, lot(date(2026, 2, 1), 90.0), 110.0, TODAY) == (pytest.approx(12.0), ASSET_TEST)
    r = Reynders(ReyndersBook({BOND: ReyndersFund(BOND)}))
    assert r.unit_base(BOND, lot(date(2026, 2, 1), 90.0), 110.0, TODAY) == (20.0, WORST_CASE)


def test_base_is_capped_at_the_real_gain_and_zero_on_a_loss():
    f = ReyndersFund(BOND, tis=[(date(2026, 1, 15), 0.0), (date(2026, 10, 1), 30.0)])
    r = Reynders(ReyndersBook({BOND: f}))
    assert r.unit_base(BOND, lot(date(2026, 2, 1), 90.0), 110.0, TODAY)[0] == 20.0
    assert r.unit_base(BOND, lot(date(2026, 2, 1), 120.0), 110.0, TODAY)[0] == 0.0
    r = Reynders(ReyndersBook({BOND: f}), ReyndersRules(cap_at_gain=False))
    assert r.unit_base(BOND, lot(date(2026, 2, 1), 120.0), 110.0, TODAY)[0] == 30.0


def test_bought_before_the_first_tis_counts_from_zero():
    f = ReyndersFund(BOND, tis=[(date(2026, 1, 15), 4.0)])
    r = Reynders(ReyndersBook({BOND: f}))
    assert r.unit_base(BOND, lot(date(2020, 1, 1), 90.0), 110.0, TODAY) == (4.0, TIS)


def test_not_subject_or_unknown_fund_pays_nothing():
    r = Reynders(ReyndersBook({BOND: ReyndersFund(BOND, subject=False)}))
    assert r.unit_base(BOND, lot(date(2026, 2, 1), 90.0), 110.0, TODAY) == (0.0, "")
    assert not r.applies("OTHER")


def test_reynders_part_comes_out_of_the_10_percent_gain():
    f = ReyndersFund(BOND, tis=[(date(2026, 1, 15), 2.0), (date(2026, 10, 1), 5.0)])
    rep = make(f)
    line = rep.tax.open_lines()[0]
    assert line.reynders_base == pytest.approx(300.0) and line.reynders_method == TIS
    assert line.taxable_gain == pytest.approx(100 * (110 - 90) - 300.0)
    s = rep.reynders_summary([line])
    assert s["tax"] == pytest.approx(90.0) and s["people"]["you"]["base"] == pytest.approx(300.0)
    keep = make(f, rules=ReyndersRules(deduct_from_cgt=False)).tax.open_lines()[0]
    assert keep.taxable_gain == pytest.approx(2000.0) and keep.reynders_base == pytest.approx(300.0)


def test_pre_2026_units_deduct_at_most_their_10_percent_gain():
    # held since 2020 at 50, 31-12-2025 at 100, now 110: the 10% gain is 10/unit, the Reynders base 40/unit (asset test)
    txns = [tx(datetime(2020, 1, 1), BUY, BOND, 10, 500)]
    line = make(ReyndersFund(BOND, debt_pct=0.666666666667), txns=txns).tax.open_lines()[0]
    assert line.reynders_base == pytest.approx(10 * 40.0, rel=1e-6)
    assert line.taxable_gain == pytest.approx(0.0, abs=1e-6)


def test_sold_lines_carry_the_reynders_base():
    txns = [tx(datetime(2026, 2, 1), BUY, BOND, 100, 9000), tx(datetime(2026, 6, 1), SELL, BOND, -40, 4400)]
    rep = make(ReyndersFund(BOND, debt_pct=0.5), txns=txns)
    sold = rep.sold_lines()[0]
    assert sold.reynders_base == pytest.approx(40 * 10.0)
    assert sold.taxable_gain == pytest.approx(40 * 20.0 - 400.0)
    assert list(rep.reynders_table().status) == ["Sold", "Open"]


def test_max_without_tax_stays_exact_with_the_deduction():
    rep = make(ReyndersFund(BOND, debt_pct=0.5))
    h = rep.portfolio.holdings()[0]
    q = rep.tax.max_without_tax(h, 500.0)          # 10% gain per unit = 20 − 10 = 10
    assert q == 50


def test_partner_share_splits_the_reynders_base():
    prof = TaxProfile()
    prof.set_partner("Partner")
    prof.household.brokers = {"B": {"you": 0.5, "partner": 0.5}}
    rep = make(ReyndersFund(BOND, debt_pct=0.5), profile=prof)
    s = rep.reynders_summary(rep.tax.open_lines())
    assert s["people"]["you"]["base"] == pytest.approx(500.0) == s["people"]["partner"]["base"]


def test_bond_etf_without_entry_is_flagged_but_not_taxed():
    rep = make(None)
    assert rep.tax.open_lines()[0].reynders_base == 0.0
    assert len(rep.reynders_flags()) == 1 and BOND in rep.reynders_flags()[0]


def test_subject_funds_are_left_out_of_harvesting():
    txns = [tx(datetime(2024, 1, 1), BUY, BOND, 100, 9000)]
    assert plan_harvest(make(None, txns=txns)).lines            # not marked: still proposed
    plan = plan_harvest(make(ReyndersFund(BOND, debt_pct=0.9), txns=txns))
    assert plan.lines == [] and any("Reynders" in n for n in plan.notes)


def test_json_round_trip(tmp_path):
    book = ReyndersBook({BOND: ReyndersFund(BOND, True, "USD", 0.95, [(date(2026, 10, 1), 5.0), (date(2026, 1, 15), 2.0)],
                                            "https://example.org/tis", "checked")})
    path = tmp_path / "reynders.json"
    book.save(path)
    again = ReyndersBook.load(path)
    assert again.to_dict() == book.to_dict()
    assert again.funds[BOND].tis_at(date(2026, 5, 1)) == (date(2026, 1, 15), 2.0)
    assert ReyndersBook.load(tmp_path / "none.json").funds == {}


def test_switch_counts_reynders_tax_in_the_cost_now():
    from meerwaarde.engine.switch import compare, sources_for, spread_plan
    from meerwaarde.funds import Fund
    rep = make(ReyndersFund(BOND, debt_pct=0.5))
    srcs = sources_for(rep, {("B", BOND)})
    target = Fund("XX0000000088", "Target", ter=0.002, accumulating=True, tob_rate=0.0012, verified=True)
    case = compare(rep, srcs, [target])[0]
    assert case.reynders_now == pytest.approx(0.30 * 100 * 10.0)
    assert case.now == pytest.approx(case.cgt_now + case.reynders_now + case.costs)
    plan = spread_plan(rep, srcs, case)
    assert plan.steps[0].costs >= 0.30 * 100 * 10.0       # all units fit this year's room; their Reynders tax is paid


def test_exports_have_the_reynders_sheet_pdf_section_and_checklist():
    import io
    from openpyxl import load_workbook
    from meerwaarde.export import tax_report_xlsx, to_xlsx
    from meerwaarde.pdf_report import tax_report_pdf
    from meerwaarde.taxreport import tax_year_report
    txns = [tx(datetime(2026, 2, 1), BUY, BOND, 100, 9000), tx(datetime(2026, 6, 1), SELL, BOND, -40, 4400)]
    rep = make(ReyndersFund(BOND, debt_pct=0.5), txns=txns)
    ws = load_workbook(io.BytesIO(to_xlsx(rep)))["Reynders tax"]
    head = [c.value for c in ws[1]]
    assert head[:3] == ["Broker", "Asset", "Name"] and ws.max_row >= 5
    formulas = [c.value for row in ws.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("=")]
    assert any("SUM(" in f for f in formulas) and any("*B" in f for f in formulas)
    ytr = tax_year_report(rep, 2026)
    assert ytr.reynders["tax"] == pytest.approx(0.30 * 400.0)
    assert any(c.startswith("Reynders tax") for c in ytr.checklist)
    assert "Reynders 2026" in load_workbook(io.BytesIO(tax_report_xlsx(ytr))).sheetnames
    assert tax_report_pdf(ytr)[:4] == b"%PDF"
