"""Excel export on a small synthetic portfolio (no real data)."""
from datetime import date, datetime

import pytest
from openpyxl import load_workbook

from meerwaarde.engine.portfolio import Portfolio
from meerwaarde.engine.tax import TaxCalculator, TaxRules
from meerwaarde.export import EUR, SHEETS, build_workbook, default_name, to_xlsx
from meerwaarde.models import BUY, DIVIDEND, SELL, Transaction
from meerwaarde.prices import PriceBook
from meerwaarde.report import Report

TODAY = date(2026, 10, 14)


def tx(when, kind, asset, qty=0.0, total=None, gross=None, cls="ETF", broker="Alpha", **kw):
    return Transaction(broker=broker, ts=when, kind=kind, asset=asset, name=f"Fund {asset}", asset_class=cls,
                       qty=qty, total_eur=total, gross_eur=gross, fees_eur=0.0, source=f"{broker.lower()}.pdf", **kw)


@pytest.fixture
def rep():
    txns = [
        tx(datetime(2024, 3, 1), BUY, "XX0000000001", 100, total=-5000, gross=5000),
        tx(datetime(2026, 5, 4), SELL, "XX0000000001", -20, total=1400, gross=1400),
        tx(datetime(2026, 2, 1), BUY, "COIN", 2.0, total=-200, gross=200, cls="Crypto", broker="Beta"),
        tx(datetime(2026, 6, 1), DIVIDEND, "XX0000000001", amount_eur=12.5),
    ]
    pf = Portfolio(txns)
    pb = PriceBook(fetch=False, today=TODAY)
    pb.set_override("XX0000000001", now=80.0, reference=60.0)
    pb.set_override("COIN", now=90.0)
    calc = TaxCalculator(pf, pb, TaxRules(), today=TODAY)
    return Report(pf, pb, calc, pf.warnings)


def load(data):
    import io
    return load_workbook(io.BytesIO(data))


def values(ws, col):
    return [ws.cell(r, col).value for r in range(2, ws.max_row + 1)]


def test_workbook_has_every_sheet_with_frozen_headers(rep):
    wb = load(to_xlsx(rep))
    assert wb.sheetnames == list(SHEETS)
    for name in ("Positions", "Sales", "Income", "Meerwaardebelasting", "Prices"):
        ws = wb[name]
        assert ws.freeze_panes == "A2", name
        assert ws.auto_filter.ref.startswith("A1:")
        assert ws["A1"].value in ("Broker", "Asset")


def test_engine_results_are_values_and_arithmetic_is_formulas(rep):
    wb = load(to_xlsx(rep))
    pos = wb["Positions"]
    head = [c.value for c in pos[1]]
    value_col = head.index("Value (€)") + 1
    assert str(pos.cell(2, value_col).value).startswith("=IF(")
    assert pos.cell(2, head.index("Cost incl. costs (€)") + 1).number_format == EUR

    sales = wb["Sales"]
    head = [c.value for c in sales[1]]
    assert sales.cell(2, head.index("Profit (€)") + 1).value == "=G2-H2"

    tax = wb["Meerwaardebelasting"]
    head = [c.value for c in tax[1]]
    gains = values(tax, head.index("Taxable gain (€)") + 1)
    expected = rep.tax_lines()
    assert sorted(gains) == pytest.approx(sorted(expected.taxable_gain_eur))
    # sold 20 units held before 2026 at 70: purchase value is the 31-12-2025 price (60) -> gain 200
    status = values(tax, head.index("Status") + 1)
    assert gains[status.index("Sold")] == pytest.approx(200.0)


def test_summary_is_formulas_over_the_sheets_with_editable_exemption(rep):
    ws = load(to_xlsx(rep))["Summary"]
    cells = {c.value: c for row in ws.iter_rows() for c in row if isinstance(c.value, str)}
    assert "Profit per broker" in cells and "Realized profit per year" in cells
    brokers = [ws.cell(r, 1).value for r in range(6, 8)]
    assert brokers == ["Alpha", "Beta"]
    assert ws.cell(6, 2).value.startswith("=SUMIFS('Sales'!")
    ex = ws.cell(cells["Exemption (change it here)"].row, 2)
    assert ex.value == 10000.0 and ex.font.color.rgb.endswith("0000FF")
    tax = ws.cell(cells["Meerwaardebelasting"].row, 2).value
    assert tax.startswith("=B") and "*B" in tax


def test_simulation_is_exported(rep):
    sim = rep.simulate(sell={("Alpha", "XX0000000001"): 10}, extra={"crypto": -50.0})
    ws = load(to_xlsx(rep, simulation=sim))["Simulator"]
    rows = [[c.value for c in r] for r in ws.iter_rows()]
    header = next(i for i, r in enumerate(rows) if r[0] == "Broker")
    first = rows[header + 1]
    assert first[1] == "XX0000000001" and first[6] == 10           # sold positions come first
    assert first[7] == pytest.approx(10 * (80 - 60))                # gain from the 31-12-2025 price
    crypto = next(r for r in rows if r[0] == "Crypto-assets")
    assert crypto[1] == -50.0
    assert sim.summary["net"] == pytest.approx(200 + 200 - 50)      # sold this year + simulated + other


def test_simulate_max_without_tax_matches_the_tax_engine(rep):
    sim = rep.simulate()
    row = sim.positions.set_index("asset").loc["XX0000000001"]
    room = 10000 - 200                                               # exemption minus this year's sale
    assert row.max_qty == rep.tax.max_without_tax(next(h for h in rep.portfolio.holdings() if h.asset == "XX0000000001"), room)


def test_writes_to_a_path_and_handles_an_empty_report(tmp_path, rep):
    p = to_xlsx(rep, tmp_path / "sub" / default_name(rep))
    assert p.name == "meerwaarde-2026-10-14.xlsx" and p.exists()

    empty = Report(Portfolio([]), PriceBook(fetch=False, today=TODAY), None, [])
    empty.tax = TaxCalculator(empty.portfolio, empty.prices, TaxRules(), today=TODAY)
    wb = build_workbook(empty)
    assert wb["Summary"]["A1"].value
    assert wb["Sales"].max_row == 1
