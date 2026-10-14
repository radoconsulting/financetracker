"""Manual positions: funds, policies, savings and gold that no export covers (synthetic data)."""
from datetime import date

import pytest

from meerwaarde.engine.harvest import plan_harvest
from meerwaarde.manual import ManualBook, ManualPosition, ManualTrade, asset_code
from meerwaarde.models import BUY, SELL
from meerwaarde.report import build_report

TODAY = date(2026, 10, 14)


def book(*positions):
    return ManualBook(list(positions))


def fund(**kw):
    d = dict(asset="XX0000000009", name="Bank fund", asset_class="Fund", broker="Bank",
             trades=[ManualTrade(date(2021, 3, 1), 10, 100.0, 2.5)], price_now=130.0, price_now_date=date(2026, 10, 1),
             price_31dec=120.0, id="p1")
    d.update(kw)
    return ManualPosition(**d)


def report(b, **kw):
    return build_report([], fetch_prices=False, today=TODAY, parsed=[], manual=b, **kw)


def test_trades_become_buy_and_sell_transactions_with_costs_apart():
    p = fund(trades=[ManualTrade(date(2021, 3, 1), 10, 100.0, 2.5), ManualTrade(date(2026, 5, 4), -4, 125.0, 1.0)])
    buy, sell = p.transactions()
    assert (buy.kind, buy.qty, buy.gross_eur, buy.fees_eur, buy.total_eur) == (BUY, 10, 1000.0, 2.5, -1002.5)
    assert (sell.kind, sell.qty, sell.gross_eur, sell.total_eur) == (SELL, -4, 500.0, 499.0)
    assert buy.broker == "Bank" and buy.source == "manual entry" and buy.ref == "p1"


def test_position_counts_in_profit_and_tax_with_the_31dec_value():
    p = fund(trades=[ManualTrade(date(2021, 3, 1), 10, 100.0, 2.5), ManualTrade(date(2026, 5, 4), -4, 125.0, 1.0)])
    rep = report(book(p))
    h = rep.holdings().iloc[0]
    assert h.qty == 6 and h.price_eur == 130.0 and h.price_source == "Your entry (manual position)"
    assert h.price_date == date(2026, 10, 1)
    assert rep.realized().pnl_eur.iloc[0] == pytest.approx(499.0 - 4 * 100.25)
    sold = rep.sold_lines()[0]
    assert sold.law_type == "fin" and sold.reference_price == 120.0
    assert sold.taxable_gain == pytest.approx(4 * (125.0 - 120.0))   # costs left out, 31-12-2025 value as purchase value


def test_insurance_and_gold_get_their_own_law_type_and_bought_in_2026_uses_the_price():
    pol = fund(asset="POLICY-1", name="Tak 23 policy", asset_class="Insurance", broker="Insurer", id="p2",
               trades=[ManualTrade(date(2026, 2, 1), 1, 20000.0, 400.0)], price_now=21000.0, price_31dec=None)
    gold = fund(asset="GOLD-1OZ", name="Gold coin", asset_class="Gold", broker="Dealer", id="p3",
                trades=[ManualTrade(date(2019, 6, 1), 2, 1200.0)], price_now=3800.0, price_31dec=3700.0)
    rep = report(book(pol, gold))
    lines = {l.asset: l for l in rep.tax.open_lines()}
    assert lines["POLICY-1"].law_type == "ins" and lines["POLICY-1"].taxable_gain == pytest.approx(1000.0)
    assert lines["GOLD-1OZ"].law_type == "gold" and lines["GOLD-1OZ"].taxable_gain == pytest.approx(200.0)


def test_eur_savings_never_give_a_gain():
    sav = fund(asset="EUR", name="Savings account", asset_class="Cash/FX", broker="Bank", id="p4",
               trades=[ManualTrade(date(2022, 1, 1), 5000, 1.0)], price_now=None, price_31dec=None)
    rep = report(book(sav))
    assert rep.holdings().value_eur.iloc[0] == 5000.0
    assert rep.tax.open_lines()[0].taxable_gain == 0.0


def test_override_wins_over_the_manual_price():
    rep = report(book(fund()), overrides_csv="asset,price_now,price_31dec\nXX0000000009,140,\n")
    h = rep.portfolio.holdings()[0]
    assert rep.prices.now(h).value == 140.0 and rep.prices.now(h).source == "Your price"
    assert rep.prices.reference(h).value == 120.0


def test_missing_31dec_value_is_flagged_not_guessed():
    rep = report(book(fund(price_31dec=None)))
    assert rep.tax.open_lines()[0].missing == "31-12-2025 price"


def test_invalid_entries_are_left_out_with_a_warning():
    bad = fund(id="p9", trades=[ManualTrade(date(2021, 1, 1), 1, 10.0), ManualTrade(date(2022, 1, 1), -3, 10.0)])
    rep = report(book(bad, fund(asset="", id="p8")))
    assert rep.portfolio.holdings() == []
    assert any("sells more than was bought" in w for w in rep.warnings)
    assert any("no asset code" in w for w in rep.warnings)


def test_ownership_per_broker_applies_to_manual_positions():
    from meerwaarde.profile import TaxProfile
    prof = TaxProfile()
    prof.set_partner("Partner")
    prof.household.brokers = {"Bank": {"you": 0.5, "partner": 0.5}}
    p = fund(trades=[ManualTrade(date(2021, 3, 1), 10, 100.0), ManualTrade(date(2026, 5, 4), -4, 125.0)])
    rep = report(book(p), profile=prof)
    people = rep.tax_summary(rep.sold_lines())["people"]
    assert people["you"]["net"] == pytest.approx(10.0) and people["partner"]["net"] == pytest.approx(10.0)


def test_insurance_is_never_proposed_for_harvesting():
    pol = fund(asset="POLICY-1", name="Tak 23 policy", asset_class="Insurance", broker="Insurer",
               trades=[ManualTrade(date(2024, 1, 1), 1, 10000.0)], price_now=13000.0, price_31dec=11000.0)
    plan = plan_harvest(report(book(pol)))
    assert plan.lines == []


def test_json_round_trip(tmp_path):
    b = book(fund(note="via the bank"), fund(asset="GOLD-1OZ", asset_class="Gold", id="p2", price_now_date=None))
    path = tmp_path / "manual.json"
    b.save(path)
    again = ManualBook.load(path)
    assert again.to_dict() == b.to_dict()
    assert again.positions[0].trades[0].costs == 2.5 and again.positions[0].price_now_date == date(2026, 10, 1)
    assert ManualBook.load(tmp_path / "missing.json").positions == []
    assert again.new_id() == "p3"


def test_asset_code_from_name():
    assert asset_code("Tak 23 Pensioen") == "TAK-23-PENSIOEN"
    assert asset_code("€€") == "MANUAL"


def test_excel_has_a_manual_positions_sheet_with_a_cash_formula():
    import io
    from openpyxl import load_workbook
    from meerwaarde.export import SHEETS, to_xlsx
    p = fund(trades=[ManualTrade(date(2021, 3, 1), 10, 100.0, 2.5), ManualTrade(date(2026, 5, 4), -4, 125.0, 1.0)])
    wb = load_workbook(io.BytesIO(to_xlsx(report(book(p)))))
    assert "Manual positions" in SHEETS
    ws = wb["Manual positions"]
    head = [c.value for c in ws[1]]
    col = head.index("Cash effect after costs (€)") + 1
    assert ws.max_row == 3
    assert ws.cell(2, col).value.startswith("=IF(")
    assert ws.cell(2, head.index("Quantity (− = sold)") + 1).value == 10
