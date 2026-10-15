"""Net worth over time (synthetic data, no network)."""
import io
from datetime import date, datetime

import pytest

from meerwaarde.engine.networth import networth
from meerwaarde.engine.portfolio import Portfolio
from meerwaarde.engine.tax import TaxCalculator, TaxRules
from meerwaarde.models import BUY, CASH_IN, DIVIDEND, REDEMPTION, SELL, Transaction
from meerwaarde.prices import PriceBook
from meerwaarde.report import Report

TODAY = date(2026, 3, 15)


def tx(when, kind, asset, broker="A", **kw):
    return Transaction(broker, datetime.combine(when, datetime.min.time()), kind, asset, name=asset, **kw)


def make(txns, now, history=None):
    pf = Portfolio(txns)
    pb = PriceBook(fetch=False, today=TODAY)
    for a, v in now.items():
        pb.set_override(a, now=v)
    history = history or {}
    pb.history = lambda asset, cls, start, end: (history.get(asset, {}), "test" if asset in history else "")
    return Report(pf, pb, TaxCalculator(pf, pb, TaxRules(), today=TODAY), pf.warnings)


def test_positions_and_cash_per_month():
    txns = [tx(date(2026, 1, 2), CASH_IN, "EUR", amount_eur=1000.0, asset_class="Cash/FX"),
            tx(date(2026, 1, 5), BUY, "X", qty=8, gross_eur=800, total_eur=-805, asset_class="ETF"),
            tx(date(2026, 2, 10), DIVIDEND, "X", amount_eur=10.0, asset_class="ETF")]
    hist = {"X": {date(2026, 1, 5): 100.0, date(2026, 1, 31): 105.0, date(2026, 2, 28): 110.0}}
    nw = networth(make(txns, {"X": 120.0}, hist), fetch=True)
    assert nw.dates == [date(2026, 1, 31), date(2026, 2, 28), date(2026, 3, 15)]
    assert nw.cash["A"] == [195.0, 205.0, 205.0]
    assert nw.by_class["Equity"] == [840.0, 880.0, 960.0]
    assert nw.total == [1035.0, 1085.0, 1165.0]


def test_missing_deposits_leave_the_cash_out():
    txns = [tx(date(2026, 1, 5), BUY, "X", qty=8, gross_eur=800, total_eur=-800, asset_class="ETF")]
    nw = networth(make(txns, {"X": 100.0}))
    assert nw.cash == {} and any("cash left out" in n for n in nw.notes)
    assert nw.total[-1] == pytest.approx(800.0)


def test_short_dip_is_kept_and_bond_payout_counts():
    txns = [tx(date(2026, 1, 30), BUY, "B", qty=1000, gross_eur=1000, total_eur=-1000, asset_class="Bond"),
            tx(date(2026, 2, 1), CASH_IN, "EUR", amount_eur=100000.0, asset_class="Cash/FX"),   # deposit after the buy
            tx(date(2026, 3, 1), SELL, "B", qty=-1000, gross_eur=0, total_eur=0, asset_class="Bond"),
            tx(date(2026, 3, 1), REDEMPTION, "B", amount_eur=1000.0, asset_class="Bond")]
    nw = networth(make(txns, {}))
    assert "A" in nw.cash and nw.cash["A"][0] == 0.0                    # −1,000 at the end of January, shown as 0
    assert nw.cash["A"][-1] == pytest.approx(100000.0)                 # the payout is cash again


def test_networth_table_and_excel_sheet():
    from openpyxl import load_workbook
    from meerwaarde.export import to_xlsx
    txns = [tx(date(2026, 1, 2), CASH_IN, "EUR", amount_eur=1000.0, asset_class="Cash/FX"),
            tx(date(2026, 1, 5), BUY, "X", qty=8, gross_eur=800, total_eur=-800, asset_class="ETF"),
            tx(date(2026, 1, 6), BUY, "BTC", broker="C", qty=1, gross_eur=100, total_eur=-100, asset_class="Crypto")]
    rep = make(txns, {"X": 100.0, "BTC": 120.0})
    t = rep.networth_table(rep.networth(), "broker")
    assert list(t.columns) == ["date", "A", "C", "total"]
    ws = load_workbook(io.BytesIO(to_xlsx(rep)))["Net worth"]
    assert ws["A1"].value == "Month end" and str(ws.cell(2, 4).value).startswith("=SUM(") and len(ws._charts) == 1
