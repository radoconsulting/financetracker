"""Fund red flags (synthetic positions, real public fund facts from the catalog)."""
import io
from datetime import date, datetime

import pytest

from meerwaarde.engine.portfolio import Portfolio
from meerwaarde.engine.redflags import alternative_for, check_funds
from meerwaarde.engine.tax import TaxCalculator, TaxRules
from meerwaarde.funds import lookup
from meerwaarde.models import BUY, DIVIDEND, TOB, Transaction
from meerwaarde.prices import PriceBook
from meerwaarde.report import Report

TODAY = date(2026, 10, 14)
IUSA, ISAC, IWDA = "IE0031442068", "IE00B6R52259", "IE00B4L5Y983"


def tx(when, kind, asset, name, cls="ETF", **kw):
    return Transaction("A", datetime.combine(when, datetime.min.time()), kind, asset, name=name, asset_class=cls, **kw)


def make(txns, now):
    pf = Portfolio(txns)
    pb = PriceBook(fetch=False, today=TODAY)
    for a, v in now.items():
        pb.set_override(a, now=v)
    return Report(pf, pb, TaxCalculator(pf, pb, TaxRules(), today=TODAY), pf.warnings)


def test_flags_and_alternatives():
    txns = [tx(date(2025, 1, 2), BUY, IUSA, "iShares Core S&P 500 (Dist)", qty=100, gross_eur=5000, total_eur=-5006),
            tx(date(2026, 3, 1), DIVIDEND, IUSA, "iShares Core S&P 500 (Dist)", amount_eur=50.0),
            tx(date(2026, 4, 1), BUY, ISAC, "iShares MSCI ACWI", qty=10, gross_eur=1000, total_eur=-1013.2),
            tx(date(2026, 4, 1), TOB, ISAC, "iShares MSCI ACWI", amount_eur=-13.2, ref="separate"),
            tx(date(2026, 4, 1), BUY, "XX0000000001", "SOME CORP BOND UCITS ETF DIST", cls="Bond ETF", qty=10,
               gross_eur=1000, total_eur=-1000),
            tx(date(2026, 4, 1), BUY, IWDA, "iShares Core MSCI World (Acc)", qty=10, gross_eur=1000, total_eur=-1000)]
    rep = make(txns, {IUSA: 60.0, ISAC: 100.0, "XX0000000001": 100.0, IWDA: 100.0})
    checks = {c.isin: c for c in check_funds(rep)}
    iusa = checks[IUSA]
    assert [f.kind for f in iusa.flags] == ["distributing"] and iusa.flags[0].per_year == pytest.approx(0.30 * 50.0)
    assert iusa.alternative.ticker == "CSPX"
    isac = checks[ISAC]
    assert any(f.kind == "tob" for f in isac.flags) and isac.alternative.index == "MSCI ACWI"
    bond = checks["XX0000000001"]
    assert {f.kind for f in bond.flags} >= {"bonds", "distributing", "ter"} and bond.alternative is None
    assert IWDA not in checks                                   # accumulating, 0.12%, cheap enough: nothing to flag


def test_alternative_is_like_for_like_only():
    alt, why = alternative_for(None, "XX0000000009", "XTRACKERS MSCI WORLD EX USA")
    assert alt is None and "no comparable" in why
    alt, _ = alternative_for(None, "XX0000000009", "Some MSCI World fund")
    assert alt is not None and alt.accumulating and alt.tob_rate <= 0.0012
    assert alternative_for(None, "XX0000000008", "SOME WORLD CORP BOND ETF", bond=True)[0] is None   # no world equity fallback


def test_red_flags_sheet():
    from openpyxl import load_workbook
    from meerwaarde.export import to_xlsx
    rep = make([tx(date(2025, 1, 2), BUY, IUSA, "iShares Core S&P 500 (Dist)", qty=10, gross_eur=500, total_eur=-500)],
               {IUSA: 60.0})
    ws = load_workbook(io.BytesIO(to_xlsx(rep)))["Red flags"]
    assert ws["B1"].value == "Fund" and ws["B2"].value.startswith("iShares Core S&P 500")
